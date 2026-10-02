"""Formatos da exportação (US1, US2, US10; spec FR-070 a FR-084; contracts/pacote-de-dados.md)."""

import io

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone

from tests.analitico import construcao as c
from tests.exportacao import leitura
from trajetoria.analitico.consultas import indicadores_do_snapshot
from trajetoria.analitico.models import SnapshotAnalitico
from trajetoria.analitico.operacoes import capturar_snapshot
from trajetoria.exportacao.contrato import (
    CABECALHO_DICIONARIO,
    CABECALHO_METADADOS,
    COLUNAS_BASE,
)
from trajetoria.exportacao.dataset import dataset_exportado
from trajetoria.exportacao.formatos import (
    csv_do_dataset,
    exportar_csv,
    exportar_xlsx,
    xlsx_do_dataset,
)
from trajetoria.exportacao.regras import ExportacaoRecusada, Motivo

pytestmark = pytest.mark.django_db

NOMES_BASE = [b.nome for b in COLUNAS_BASE]


def _metadados(tabelas) -> dict:
    return {chave: valor for chave, valor, _ in tabelas["metadados"][1:]}


# --- US1: pacote CSV ------------------------------------------------------------------------------


def test_pacote_csv_com_exatamente_tres_arquivos_utf8_sem_bom_crlf(snapshot):
    entradas = leitura.entradas_csv(exportar_csv(snapshot))
    for nome, conteudo in entradas.items():
        assert not conteudo.startswith(b"\xef\xbb\xbf"), nome
        conteudo.decode("utf-8")
        assert conteudo.endswith(b"\r\n"), nome
        assert b"\n" not in conteudo.replace(b"\r\n", b""), nome  # nenhuma quebra solta


def test_cabecalhos_do_pacote_csv(snapshot):
    tabelas = leitura.ler_csv(exportar_csv(snapshot))
    assert tabelas["dados"][0][: len(NOMES_BASE)] == NOMES_BASE
    assert tuple(tabelas["dicionario"][0]) == CABECALHO_DICIONARIO
    assert tuple(tabelas["metadados"][0]) == CABECALHO_METADADOS
    colunas_no_dicionario = [
        linha[1] for linha in tabelas["dicionario"][1:] if linha[0] == "coluna"
    ]
    assert colunas_no_dicionario == tabelas["dados"][0]


def test_metadados_identificam_o_snapshot_e_contam_os_registros(snapshot):
    tabelas = leitura.ler_csv(exportar_csv(snapshot))
    metadados = _metadados(tabelas)
    assert metadados["snapshot"] == str(snapshot.pk)
    assert metadados["registros"] == str(indicadores_do_snapshot(snapshot).registros)
    assert len(tabelas["dados"]) - 1 == indicadores_do_snapshot(snapshot).registros
    assert metadados["capturado_em"] == timezone.localtime(snapshot.capturado_em).isoformat()


def test_booleanos_e_ausencia_no_csv(snapshot):
    tabelas = leitura.ler_csv(exportar_csv(snapshot))
    cabecalho = tabelas["dados"][0]
    elegivel = cabecalho.index("elegivel_no_snapshot")
    assert {linha[elegivel] for linha in tabelas["dados"][1:]} <= {"true", "false"}
    bruto = leitura.entradas_csv(exportar_csv(snapshot))["dados.csv"]
    assert b'""' not in bruto  # ausência é campo vazio, nunca texto vazio entre aspas


def test_exportar_sem_snapshot_ou_com_campanha(cenario):
    with pytest.raises(TypeError):
        exportar_csv()  # type: ignore[call-arg]
    with pytest.raises(TypeError):
        exportar_csv(cenario.campanha)
    with pytest.raises(TypeError):
        dataset_exportado(cenario.campanha)


def test_snapshot_nao_gravado(cenario):
    with pytest.raises(ExportacaoRecusada) as erro:
        exportar_csv(SnapshotAnalitico(campanha=cenario.campanha))
    assert erro.value.motivo is Motivo.SNAPSHOT_NAO_GRAVADO


def test_sem_chave_recusa_antes_de_ler_registros(snapshot, settings):
    settings.TRAJETORIA_CHAVE_PSEUDONIMIZACAO = ""
    with CaptureQueriesContext(connection) as consultas, pytest.raises(ExportacaoRecusada) as erro:
        exportar_csv(snapshot)
    assert erro.value.motivo is Motivo.CHAVE_AUSENTE
    assert not [q for q in consultas.captured_queries if "registrodosnapshot" in q["sql"]]


# --- US2: XLSX com a mesma semântica --------------------------------------------------------------


def test_xlsx_com_tres_abas_e_cabecalhos(snapshot):
    tabelas = leitura.ler_xlsx(exportar_xlsx(snapshot))
    assert tabelas["dados"][0][: len(NOMES_BASE)] == NOMES_BASE
    assert tuple(tabelas["dicionario"][0]) == CABECALHO_DICIONARIO
    assert tuple(tabelas["metadados"][0]) == CABECALHO_METADADOS


def test_xlsx_tipos_nativos_e_ausencia(snapshot):
    tabelas = leitura.ler_xlsx(exportar_xlsx(snapshot))
    cabecalho = tabelas["dados"][0]
    elegivel, ano = cabecalho.index("elegivel_no_snapshot"), cabecalho.index("ano_conclusao")
    unidade = cabecalho.index("unidade")
    linhas = tabelas["dados"][1:]
    assert all(isinstance(linha[elegivel], bool) for linha in linhas)
    assert {type(linha[ano]) for linha in linhas} <= {int, type(None)}
    assert None in {linha[ano] for linha in linhas}  # vitoria_sem_atributos
    assert "" not in {linha[unidade] for linha in linhas}


def test_xlsx_sem_formulas_nem_recursos_de_bi(snapshot):
    from openpyxl import load_workbook

    conteudo = exportar_xlsx(snapshot)
    assert leitura.formulas(conteudo) == []
    pasta = load_workbook(io.BytesIO(conteudo))
    for planilha in pasta.worksheets:
        assert planilha.sheet_state == "visible"
        assert not planilha.merged_cells.ranges
        assert planilha.auto_filter.ref is None
        assert not planilha.conditional_formatting
        assert not planilha.data_validations.dataValidation
        assert not planilha._charts and not planilha._images


def test_csv_xlsx_e_dataset_logico_sao_identicos(snapshot):
    logico = leitura.normalizar_dataset(dataset_exportado(snapshot))
    csv_lido = leitura.normalizar_csv(leitura.ler_csv(exportar_csv(snapshot)))
    xlsx_lido = leitura.normalizar_xlsx(leitura.ler_xlsx(exportar_xlsx(snapshot)))
    for tabela in ("dados", "dicionario", "metadados"):
        assert csv_lido[tabela] == logico[tabela], tabela
        assert xlsx_lido[tabela] == logico[tabela], tabela


# --- US10: ordem determinística e reprodutibilidade -----------------------------------------------


def test_ordem_estrutural_secao_pergunta_opcao(inst):  # caso T
    from trajetoria.instrumento import operacoes as op

    versao = op.criar_versao(op.criar_pesquisa("Pesquisa de ordem"), "2027")
    s1, s2 = op.adicionar_secao(versao, 1), op.adicionar_secao(versao, 2)
    z = op.adicionar_pergunta(s1, 1, "TEXTO_CURTO", "Zebra", obrigatoria=False)
    a = op.adicionar_pergunta(s1, 2, "ESCOLHA_MULTIPLA", "Abacate", obrigatoria=False)
    o_z, o_a = op.adicionar_opcao(a, 1, "Zeta"), op.adicionar_opcao(a, 2, "Alfa")
    m = op.adicionar_pergunta(s2, 1, "TEXTO_CURTO", "Meio", obrigatoria=False)
    op.reordenar_secoes(versao, [s2, s1])  # s2 vira a primeira
    op.reordenar_perguntas(s1, [a, z])
    op.reordenar_opcoes(a, [o_a, o_z])
    op.publicar(versao)
    snapshot = capturar_snapshot(c.campanha_aberta_no_passado(versao))
    cabecalho = dataset_exportado(snapshot).dados.cabecalho
    perguntas = [n for n in cabecalho if n.startswith("pergunta_")]
    esperado = [
        f"pergunta_{m.pk.hex}__aplicavel",
        f"pergunta_{m.pk.hex}",
        f"pergunta_{a.pk.hex}__aplicavel",
        f"pergunta_{a.pk.hex}__opcao_{o_a.pk.hex}",
        f"pergunta_{a.pk.hex}__opcao_{o_z.pk.hex}",
        f"pergunta_{z.pk.hex}__aplicavel",
        f"pergunta_{z.pk.hex}",
    ]
    assert perguntas == esperado


def test_linhas_na_ordem_da_012(snapshot, chave_ficticia):
    from trajetoria.exportacao.pseudonimos import pseudonimo

    ordem = sorted(snapshot.registros.values_list("conclusao_id", flat=True))
    esperados = [pseudonimo("conclusao", pk, chave_ficticia) for pk in ordem]
    linhas = dataset_exportado(snapshot).dados.linhas
    assert [linha[0] for linha in linhas] == esperados


def test_csv_reproduzivel_byte_a_byte(snapshot):
    assert exportar_csv(snapshot) == exportar_csv(snapshot)


def test_xlsx_reproduzivel_em_conteudo(snapshot):
    a = leitura.normalizar_xlsx(leitura.ler_xlsx(exportar_xlsx(snapshot)))
    b = leitura.normalizar_xlsx(leitura.ler_xlsx(exportar_xlsx(snapshot)))
    assert a == b


def test_xlsx_sem_o_momento_da_exportacao(snapshot):
    from datetime import UTC, timedelta

    from openpyxl import load_workbook

    # snapshot "antigo": o arquivo deve carregar o momento da captura, não o da exportação
    SnapshotAnalitico.objects.filter(pk=snapshot.pk).update(
        capturado_em=snapshot.capturado_em - timedelta(days=3)
    )
    snapshot.refresh_from_db()
    criado = load_workbook(io.BytesIO(exportar_xlsx(snapshot))).properties.created
    capturado = snapshot.capturado_em.astimezone(UTC).replace(tzinfo=None, microsecond=0)
    assert criado.replace(tzinfo=None, microsecond=0) == capturado


def test_outra_chave_muda_so_os_pseudonimos(snapshot, settings):
    antes = dataset_exportado(snapshot)
    settings.TRAJETORIA_CHAVE_PSEUDONIMIZACAO = "outra-chave-ficticia-de-teste-0000000000002"
    depois = dataset_exportado(snapshot)
    assert antes.dicionario == depois.dicionario and antes.metadados == depois.metadados
    for x, y in zip(antes.dados.linhas, depois.dados.linhas, strict=True):
        assert x[2:] == y[2:]
        assert x[0] != y[0] and x[1] != y[1]


# --- Revisão de código: reuso e leitura única ----------------------------------------------


def test_dois_formatos_de_um_dataset_ja_montado(snapshot):
    montado = dataset_exportado(snapshot)
    assert csv_do_dataset(montado) == exportar_csv(snapshot)
    assert leitura.normalizar_xlsx(leitura.ler_xlsx(xlsx_do_dataset(montado))) == (
        leitura.normalizar_xlsx(leitura.ler_xlsx(exportar_xlsx(snapshot)))
    )


def test_conteudo_da_versao_lido_uma_vez(snapshot):
    with CaptureQueriesContext(connection) as consultas:
        dataset_exportado(snapshot)
    secoes = [q for q in consultas.captured_queries if 'FROM "instrumento_secao"' in q["sql"]]
    assert len(secoes) == 1


def test_capturado_em_do_dataset(snapshot):
    montado = dataset_exportado(snapshot)
    assert montado.capturado_em == timezone.localtime(snapshot.capturado_em)
    metadados = {chave: valor for chave, valor, _ in montado.metadados.linhas}
    assert montado.capturado_em == metadados["capturado_em"]
