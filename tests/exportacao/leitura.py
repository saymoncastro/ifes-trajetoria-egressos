"""Leitores independentes dos arquivos exportados (Feature 013; não são testes).

O CSV é lido com `zipfile` + `csv` e a reversão universal do escape (contracts/
pacote-de-dados.md); o XLSX, com openpyxl — biblioteca diferente da que escreve, para que
um erro simétrico do escritor não se esconda (research R11). As funções `normalizar_*`
levam o que foi lido à forma comparável do dataset lógico: `None`, `bool`, `int`, `date` e
`str`, com momentos como texto ISO 8601.
"""

import csv
import io
import re
import zipfile
from datetime import date, datetime

ARQUIVOS_CSV = ("dados.csv", "dicionario.csv", "metadados.csv")
ABAS_XLSX = ("Dados", "Dicionário", "Metadados")
TABELAS = ("dados", "dicionario", "metadados")


def reverter(campo: str) -> str:
    """A reversão universal do contrato: remove um apóstrofo inicial, se houver."""
    return campo[1:] if campo.startswith("'") else campo


def entradas_csv(pacote: bytes) -> dict[str, bytes]:
    """Os bytes brutos de cada CSV do pacote, exigindo exatamente os três, nesta ordem."""
    with zipfile.ZipFile(io.BytesIO(pacote)) as arquivo:
        assert arquivo.namelist() == list(ARQUIVOS_CSV)
        return {nome: arquivo.read(nome) for nome in ARQUIVOS_CSV}


def ler_csv(pacote: bytes, *, revertido: bool = True) -> dict[str, list[list[str]]]:
    """`{tabela: [cabeçalho, *linhas]}`, com cada campo já revertido (ou bruto)."""
    tabelas = {}
    for tabela, (nome, conteudo) in zip(TABELAS, entradas_csv(pacote).items(), strict=True):
        assert not conteudo.startswith(b"\xef\xbb\xbf"), f"{nome} com BOM"
        texto = conteudo.decode("utf-8")
        linhas = list(csv.reader(io.StringIO(texto, newline="")))
        tabelas[tabela] = [[reverter(c) for c in linha] if revertido else linha for linha in linhas]
    return tabelas


def ler_xlsx(conteudo: bytes) -> dict[str, list[list]]:
    """`{tabela: [cabeçalho, *linhas]}` com os valores como o openpyxl os lê."""
    from openpyxl import load_workbook

    pasta = load_workbook(io.BytesIO(conteudo))
    assert pasta.sheetnames == list(ABAS_XLSX)
    tabelas = {}
    for tabela, aba in zip(TABELAS, ABAS_XLSX, strict=True):
        planilha = pasta[aba]
        largura = len(next(planilha.iter_rows(min_row=1, max_row=1, values_only=True)))
        tabelas[tabela] = [
            list(linha) + [None] * (largura - len(linha))
            for linha in planilha.iter_rows(values_only=True)
        ]
    return tabelas


def formulas(conteudo: bytes) -> list[str]:
    """Coordenadas das células de fórmula em todas as abas (deve ser sempre vazio)."""
    from openpyxl import load_workbook

    pasta = load_workbook(io.BytesIO(conteudo))
    return [
        f"{planilha.title}!{celula.coordinate}"
        for planilha in pasta.worksheets
        for linha in planilha.iter_rows()
        for celula in linha
        if celula.data_type == "f"
    ]


# --- Normalização ----------------------------------------------------------------------------


def _do_texto(valor: str, tipo: str):
    if valor == "":
        return None
    if tipo == "booleano":
        assert valor in ("true", "false"), valor
        return valor == "true"
    if tipo == "inteiro":
        return int(valor)
    if tipo == "data":
        return date.fromisoformat(valor)
    return valor  # texto e momento (ISO)


def _tipos_do_dicionario():
    from trajetoria.exportacao.contrato import COLUNAS_DICIONARIO

    return [tipo for _, tipo in COLUNAS_DICIONARIO]


def normalizar_csv(tabelas: dict[str, list[list[str]]]) -> dict[str, tuple]:
    """Converte o CSV lido (já revertido) pelos tipos que o próprio pacote declara:
    Dicionário pelos tipos fixos do contrato; Dados pelo `tipo_valor` das linhas `coluna`
    do Dicionário; Metadados pela coluna `tipo` de cada linha."""
    dicionario = tabelas["dicionario"]
    tipos_dic = _tipos_do_dicionario()
    linhas_dic = [
        tuple(_do_texto(v, t) for v, t in zip(linha, tipos_dic, strict=True))
        for linha in dicionario[1:]
    ]
    nome, tipo = dicionario[0].index("coluna"), dicionario[0].index("tipo_valor")
    elemento = dicionario[0].index("elemento")
    tipo_da_coluna = {
        linha[nome]: linha[tipo] for linha in linhas_dic if linha[elemento] == "coluna"
    }
    dados = tabelas["dados"]
    tipos_dados = [tipo_da_coluna[c] for c in dados[0]]
    linhas_dados = [
        tuple(_do_texto(v, t) for v, t in zip(linha, tipos_dados, strict=True))
        for linha in dados[1:]
    ]
    linhas_meta = [
        (chave, _do_texto(valor, tipo), tipo) for chave, valor, tipo in tabelas["metadados"][1:]
    ]
    return {
        "dados": (tuple(dados[0]), linhas_dados),
        "dicionario": (tuple(dicionario[0]), linhas_dic),
        "metadados": (tuple(tabelas["metadados"][0]), linhas_meta),
    }


_ESCAPE_OOXML = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _da_planilha(valor):
    if isinstance(valor, str):
        # Decodificação OOXML de `_xHHHH_`, como fazem Excel e LibreOffice (o openpyxl não a
        # aplica a todos os caracteres). Inequívoca: texto com o padrão literal é recusado
        # pela exportação (research R12).
        return _ESCAPE_OOXML.sub(lambda m: chr(int(m.group(1), 16)), valor)
    if isinstance(valor, datetime):
        assert valor.time() == datetime.min.time(), valor
        return valor.date()
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor


def normalizar_xlsx(tabelas: dict[str, list[list]]) -> dict[str, tuple]:
    return {
        nome: (
            tuple(linhas[0]),
            [tuple(_da_planilha(v) for v in linha) for linha in linhas[1:]],
        )
        for nome, linhas in tabelas.items()
    }


def _logico(valor):
    return valor.isoformat() if isinstance(valor, datetime) else valor


def normalizar_dataset(dataset) -> dict[str, tuple]:
    """O dataset lógico na mesma forma: momentos como texto ISO."""
    return {
        nome: (
            tuple(tabela.cabecalho),
            [tuple(_logico(v) for v in linha) for linha in tabela.linhas],
        )
        for nome, tabela in (
            ("dados", dataset.dados),
            ("dicionario", dataset.dicionario),
            ("metadados", dataset.metadados),
        )
    }
